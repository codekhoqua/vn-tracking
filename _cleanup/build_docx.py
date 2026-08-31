import time
import io
import requests
from docx import Document
from docx.shared import Inches
from playwright.sync_api import sync_playwright

url = "https://linkstory-co.notion.site/PSD-3a1bdcae778b80168827d0747a9c4707?pvs=143"

text_content = [
    """PSDデータの新ルール
Luật mới về dữ liệu PSD

PSDをZIP圧縮し直接Mikanへアップロードできるようになりました。
Đã có thể nén PSD thành file ZIP và upload trực tiếp lên Mikan.

従来のjpgでアップロードをする方法だと、inpainting情報が削除されていましたが
Với phương pháp upload bằng jpg trước đây, thông tin inpainting đã bị xóa, nhưng

ZIPでアップロードすることでMikanにinpainting情報が残せるようになりました。
bằng cách upload bằng ZIP, thông tin inpainting hiện có thể được giữ lại trên Mikan.

さらにPhotoshopのレイヤー情報もMikan上に表示されるようになり、Mikan上のローラー・消しゴムアイコンから
Hơn nữa, thông tin layer của Photoshop cũng sẽ được hiển thị trên Mikan, và từ biểu tượng con lăn / cục tẩy trên Mikan

編集ができるので、簡単な修正はMikan上で編集可能です。
bạn có thể chỉnh sửa, do đó các sửa đổi đơn giản có thể được thực hiện ngay trên Mikan.

しかし、画像写植の文字がピクセル化されるため文字情報が残りません。さらに透明効果が入った特殊加工をしたデ
Tuy nhiên, do chữ trong typeset hình ảnh bị pixel hóa nên thông tin chữ không được giữ lại. Thêm vào đó, đối với dữ liệu

ータは見え方が変わってしまいます。
đã được xử lý đặc biệt có hiệu ứng trong suốt thì cách hiển thị sẽ bị thay đổi.

そういったページは従来通りの手順のまま、jpgアップになります。
Những trang như vậy sẽ vẫn làm theo quy trình như trước đây, tức là upload bằng jpg.

【目的】
【Mục đích】

変更となった新ルールを理解し、ルール通りに作業を進められるようにする
Hiểu rõ các luật mới đã được thay đổi, để có thể tiến hành công việc đúng theo quy tắc.

【レタッチ・写植・修正写植工程の変更】
【Thay đổi trong công đoạn Retouch / Typeset / Sửa Typeset】

●各作業者がMikanからPSDをダウンロードする
●Mỗi người làm việc sẽ download PSD từ Mikan.

※編集ページのみ
※Chỉ trang chỉnh sửa

※Dropboxからダウンロードする場合もある
※Cũng có trường hợp download từ Dropbox

●PSDのレイヤー構造
●Cấu trúc layer của PSD

●Dropboxへ格納
●Lưu trữ vào Dropbox

※フォルダ構造変更
※Thay đổi cấu trúc thư mục

●Mikanへアップロード
●Upload lên Mikan

①ZIPはPhotoshopのプラグインを使用して書き出す
①ZIP sẽ được xuất ra bằng cách sử dụng plugin của Photoshop

Orangeさんのレタッチのマニュアルに「プラグインのインストール手順」あり
Trong manual hướng dẫn retouch của Orange có "Trình tự cài đặt plugin"

・Lettering手順書…
・Tài liệu hướng dẫn Lettering...

docs.google.com
docs.google.com (Giữ nguyên)

・retouch手順書…
・Tài liệu hướng dẫn retouch...

docs.google.com
docs.google.com (Giữ nguyên)

★プラグイン置き場
★Nơi để plugin

drive.google.com
drive.google.com (Giữ nguyên)

※ZIPは必ずプラグインから書き出す
※ZIP bắt buộc phải xuất ra từ plugin""",

    """②Mikanへアップロードさせる方法が２つになる
②Sẽ có 2 cách để upload lên Mikan

・ZIP…直接Mikanへアップロード
・ZIP... Upload trực tiếp lên Mikan

　　　　（シンプルなレタッチのみ）
　　　　(Chỉ dành cho retouch đơn giản)

・JPG…従来通りのアップロード
・JPG... Upload theo cách trước đây

　　　　（画像写植・透過があるもの）
　　　　(Những cái có typeset hình ảnh / có độ trong suốt)

※JPGの手順は省略
※Bỏ qua các bước đối với JPG

Download PSD………………………PSDをダウンロード
Download PSD………………………Download PSD

Import Photoshop Inpainting…ZIPをアップロード
Import Photoshop Inpainting…Upload ZIP

インストールしたPhotoshopのプ
Nhấn vào plugin Photoshop

ラグインをクリックしてZIP保存。
đã cài đặt để lưu dưới dạng ZIP.

任意の場所に保存される。
Sẽ được lưu ở vị trí bất kỳ.""",

    """③PSDレイヤーのルール変更
③Thay đổi luật về layer PSD

（JPGでアップロードする際は従来のルール通り）
(Khi upload bằng JPG thì làm theo luật như trước đây)

アップ方法を変更する際はレイヤーを見直す必要がある
Khi thay đổi phương pháp upload, cần phải kiểm tra lại các layer

ZIPでアップロードする時はinpaintingフォルダの上に新規レイヤーを追加して作業する
Khi upload bằng ZIP, hãy thêm một layer mới lên trên thư mục inpainting để làm việc

• ZIPファイルの中には表示しているレイヤー情報のみが入っているため、非表示にしているレイヤーは削除
• Vì trong file ZIP chỉ chứa thông tin của các layer đang hiển thị, nên các layer bị ẩn sẽ bị xóa

される
đi

• レイヤーの中は、近くの修正をまとめる。
• Trong layer, hãy gom các sửa đổi ở gần nhau lại.

広域を一つのレイヤーにまとめると、Mikan上でinpainting枠が重なり複雑化するため
Vì nếu gom một khu vực rộng vào một layer, khung inpainting trên Mikan sẽ bị chồng chéo và trở nên phức tạp.

• 同じ名称レイヤーがあるとZIP保存ができないため、名称を変更する
• Nếu có layer trùng tên thì sẽ không thể lưu ZIP được, do đó hãy đổi tên.

• inpaintingファイルの中に新規レイヤーを入れない
• Không đưa layer mới vào bên trong file inpainting

既存レイヤー(Mikanですでに制作済みのレイヤー)に直接描き込むと、
Nếu vẽ trực tiếp lên layer hiện có (layer đã được tạo sẵn trên Mikan),

　　変更が検出されずに書き出し対象にならない
　　thay đổi sẽ không được phát hiện và sẽ không trở thành đối tượng được xuất ra.

ZIPの場合はinpaintingを複製しないが、JPGの場合は複製
Đối với ZIP thì không nhân bản inpainting, nhưng đối với JPG thì có nhân bản

する（従来のルール通り）
(Làm theo luật như trước đây)

ZIPでアップするデータとは
Dữ liệu sẽ được upload bằng ZIP là gì

単色で色の変化がなく特殊な処理をしていない（シンプルなデータのみ）
Chỉ có một màu, không có sự thay đổi màu sắc và không có xử lý đặc biệt (chỉ dữ liệu đơn giản)""",

    """JPGでアップするデータとは
Dữ liệu sẽ được upload bằng JPG là gì

①表紙
①Trang bìa

②扉絵(口絵)
②Trang minh họa (Trang đầu)

③奥付
③Trang colophon (Trang thông tin xuất bản)

⑤プラグインで保存ができないもの
⑤Những thứ không thể lưu bằng plugin

※エラーが出て保存ができないものがたまにある
※Thỉnh thoảng có những file bị lỗi và không thể lưu được

⑥特殊な処理をしている
⑥Đang có xử lý đặc biệt

〈特殊な処理とは〉一例
〈Xử lý đặc biệt là gì〉 Một số ví dụ

・乗算などのレイヤー効果
・Hiệu ứng layer như Multiply (Nhân bản)

（文字を透けさせる）
(Làm cho chữ trở nên trong suốt)

・クリッピングマスク
・Clipping mask (Mặt nạ cắt)

（トーンを文字の形に抜く）
(Cắt tone theo hình dạng của chữ)

・調整レイヤー
・Layer điều chỉnh

（色調補正）
(Chỉnh sửa tông màu)

・FXの特殊効果
・Hiệu ứng đặc biệt của FX

（フチぼかし・ドロップシャドウ）
(Làm mờ viền / Bóng đổ (Drop shadow))

※特に透明効果があるもの
※Đặc biệt là những thứ có hiệu ứng trong suốt

※通常マスクは問題なし
※Mask thông thường thì không sao

※幕間のロゴが単色ではないもの
※Logo ở phần chuyển cảnh (intermission) không phải là màu đơn sắc

〈理由〉
〈Lý do〉

Mikanにzipでアップをするとモノクロ２階調となり透明部分の再現ができなくなるため、透過の見え
Vì khi upload lên Mikan bằng zip, nó sẽ trở thành 2 tông màu đơn sắc (trắng đen) và không thể tái tạo lại phần trong suốt, do đó

方が変わってしまいます。
cách hiển thị độ trong suốt sẽ bị thay đổi.

特殊効果を使用している場合はjpgで保存。
Trường hợp có sử dụng hiệu ứng đặc biệt thì lưu bằng jpg.

※JPGでアップするとinpainting情報は消える
※Khi upload bằng JPG thì thông tin inpainting sẽ bị biến mất

画像写植したPSDデータはすべて文字を編集できる状態のPSDでDropboxへ格納す
Tất cả dữ liệu PSD đã typeset hình ảnh đều phải được lưu trữ vào Dropbox dưới dạng PSD ở trạng thái có thể chỉnh sửa được chữ.

る
(Phần kết câu của từ 格納する ở dòng trên)

ZIPで圧縮されているデータには文字情報は含まれていないため
Vì dữ liệu đã được nén bằng ZIP không bao gồm thông tin chữ

• ZIPでアップするデータ
• Dữ liệu upload bằng ZIP

単色で色の変化がなく特殊な処理をしていない（シンプルなデータのみ）
Chỉ có một màu, không có sự thay đổi màu sắc và không có xử lý đặc biệt (chỉ dữ liệu đơn giản)

• JPGでアップするデータ
• Dữ liệu upload bằng JPG

特殊な処理をしている
Đang có xử lý đặc biệt""",

    """④Dropboxの格納の方法について
④Về cách thức lưu trữ trên Dropbox

・編集したPSDのみを格納する
・Chỉ lưu trữ các PSD đã chỉnh sửa

※編集できる生データの状態で格納
※Lưu trữ ở trạng thái dữ liệu thô (raw data) có thể chỉnh sửa được

※全ページ格納しない
※Không lưu trữ tất cả các trang

・JPGアップロードの場合はJPGも格納する
・Trường hợp upload bằng JPG thì cũng lưu trữ cả JPG

Dropboxを参考　
Tham khảo Dropbox

02_Orange localized data＞99_テンプレ
02_Orange localized data > 99_Template

※表示されていない場合は、Dropboxの権限付与を進行係へ申請してください
※Trường hợp không hiển thị, vui lòng yêu cầu người phụ trách tiến độ cấp quyền truy cập Dropbox

JPGでアップロードする場合はフォルダを制作し格納する
Trường hợp upload bằng JPG thì tạo thư mục và lưu trữ vào đó

⑤JPGに置き換えられているページを編集する時の注意点
⑤Những lưu ý khi chỉnh sửa trang đã được thay thế bằng JPG

JPGでアップするとinpainting情報が消えるため、Mikan上で編集がしにくくなります。
Vì khi upload bằng JPG thông tin inpainting sẽ bị mất, nên sẽ khó chỉnh sửa trên Mikan.

文字情報が無くなっていたり、先祖返りをする危険があるため、
Vì có nguy cơ bị mất thông tin chữ, hoặc bị hoàn tác lại bản cũ (regression), do đó

一度JPGでアップされたページはMikan上で編集しません。
Trang đã từng được upload bằng JPG một lần thì sẽ không chỉnh sửa trên Mikan nữa.

Dropboxに格納されているPSDから編集します。
Sẽ chỉnh sửa từ PSD đang được lưu trữ trên Dropbox.

※従来通りDropboxからPSDを開き編集
※Mở PSD từ Dropbox và chỉnh sửa như trước đây

JPGアップのページはグレーの枠がついている
Trang upload bằng JPG có viền màu xám

※データが格納されていない場合は速やかにレタッチ担当者・もしくは進行係へご連絡してください。
※Nếu dữ liệu chưa được lưu trữ, vui lòng liên hệ ngay với người phụ trách retouch hoặc người phụ trách tiến độ.

※Dropboxに格納されているPSDは生データなので文字や透明効果の再編集が可能。
※PSD được lưu trữ trên Dropbox là dữ liệu thô nên có thể chỉnh sửa lại chữ và hiệu ứng trong suốt.

　こちらを使用せず間違えてMikanからZIPをダウンロードし編集してしまうと、
　Nếu không sử dụng file này mà lầm tưởng download ZIP từ Mikan về để chỉnh sửa,

　編集ができない状態で進むことになり、修正が必要になった時に時間がかかる。
　thì sẽ tiến hành trong tình trạng không thể chỉnh sửa được, và sẽ mất rất nhiều thời gian khi cần phải sửa đổi.

　グレー枠が付いているページは必ずDropboxのPSDから編集を行う。
　Đối với các trang có viền màu xám, bắt buộc phải chỉnh sửa từ file PSD trên Dropbox."""
]

def run():
    doc = Document()
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.goto(url)
        print("Waiting for page to load...")
        time.sleep(5)
        
        # Scroll down to load all lazy images
        for i in range(10):
            page.mouse.wheel(0, 500)
            time.sleep(0.5)
        
        print("Extracting images...")
        # Get all images
        images = page.locator("img").all()
        img_urls = []
        for img in images:
            src = img.get_attribute("src")
            if src:
                # Resolve relative URL if needed
                if src.startswith("/"):
                    src = "https://linkstory-co.notion.site" + src
                # Ignore small icons, only want the large content images
                box = img.bounding_box()
                if box and box["width"] > 200:
                    img_urls.append(src)
        
        print(f"Found {len(img_urls)} large images.")
        
        # Build doc
        for i, text_block in enumerate(text_content):
            # Insert image first (since they are screenshots at the top of the OCR pages)
            if i < len(img_urls):
                try:
                    response = requests.get(img_urls[i])
                    image_stream = io.BytesIO(response.content)
                    doc.add_picture(image_stream, width=Inches(6.0))
                except Exception as e:
                    print(f"Failed to add image {i}: {e}")
            
            # Insert text
            doc.add_paragraph(text_block)
            
            if i < len(text_content) - 1:
                doc.add_page_break()
                
        doc.save("PSD_New_Rules.docx")
        print("Saved to PSD_New_Rules.docx")
        browser.close()

if __name__ == "__main__":
    run()
